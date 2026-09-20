"""Tests RAG: chunking, parsers, embeddings, discovery y endpoints (sin DB)."""
from __future__ import annotations

import io
import math

import pytest
from httpx import ASGITransport, AsyncClient

from app.main import app
from app.api.v1.endpoints import rag as rag_ep
from app.rag.chunking import chunk_text, split_sections
from app.rag.embeddings import HashEmbedder
from app.rag.parsers import detect_kind, extract_text
from app.rag.sources import discover_campaigns, discover_diarios
from app.rag.store import SearchHit


# -- chunking -----------------------------------------------------------
def test_split_sections_detects_headings():
    text = '# Calidad del aire\nAlgo de texto.\n\n## PM2.5\nMas texto.'
    secs = split_sections(text)
    assert [s for s, _ in secs] == ['Calidad del aire', 'PM2.5']


def test_chunk_text_deterministic_and_bounded():
    text = '# T\n' + 'La inversion termica atrapa contaminantes en el valle. ' * 60
    a = chunk_text(text)
    b = chunk_text(text)
    assert a == b
    assert len(a) >= 2
    assert all(len(c.text) <= 1000 for c in a)
    assert all(c.section == 'T' for c in a)


def test_chunk_overlap_keeps_context():
    text = 'Tema. ' + 'Frase intermedia sobre PM2.5 en Itagui. ' * 40 + 'Cierre final.'
    chunks = chunk_text(text, target=200, overlap=60)
    assert len(chunks) >= 2
    assert 'PM2.5' in chunks[1].text


# -- parsers ------------------------------------------------------------
def test_detect_kind():
    assert detect_kind('a.PDF') == 'pdf'
    assert detect_kind('plan.xlsx') == 'excel'
    assert detect_kind('tesis.docx') == 'word'
    assert detect_kind('nota.md') == 'text'


def test_extract_txt_and_csv():
    text, meta = extract_text('n.txt', 'Hola valle'.encode())
    assert 'Hola valle' in text
    text, meta = extract_text('d.csv', 'a,b\n1,2\n'.encode())
    assert meta['format'] == 'csv' and ' | ' in text


def _mini_pdf(body: str) -> bytes:
    """PDF minimo valido con tabla xref real."""
    stream = (
        'BT /F1 24 Tf 72 720 Td (' + body + ') Tj ET'
    ).encode('latin-1')
    objs = [
        b'1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj',
        b'2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj',
        b'3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R'
        b'/Resources<</Font<</F1 5 0 R>>>>>>endobj',
        b'4 0 obj<</Length ' + str(len(stream)).encode() + b'>>stream' + chr(10).encode() + stream + chr(10).encode() + b'endstream endobj',
        b'5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj',
    ]
    out = bytearray(b'%PDF-1.4' + chr(10).encode())
    offsets = [0]
    for i, body_bytes in enumerate(objs, start=1):
        offsets.append(len(out))
        out += body_bytes + chr(10).encode()
    xref_pos = len(out)
    out += f'xref{chr(10)}0 {len(objs) + 1}{chr(10)}'.encode()
    out += f'0000000000 65535 f {chr(10)}'.encode()
    for off in offsets[1:]:
        out += f'{off:010d} 00000 n {chr(10)}'.encode()
    out += (
        f'trailer{chr(10)}<</Size {len(objs) + 1}/Root 1 0 R>>{chr(10)}'
        f'startxref{chr(10)}{xref_pos}{chr(10)}%%EOF{chr(10)}'
    ).encode()
    return bytes(out)


def test_extract_pdf_real_text():
    text, meta = extract_text('informe.pdf', _mini_pdf('Calidad del aire PM2.5 38'))
    assert meta['format'] == 'pdf'
    assert 'PM2.5' in text and '[pag 1]' in text


def test_extract_html_strips_tags():
    html = '<html><head><style>.x{color:red}</style><script>var a=1;</script></head><body><h1>Resolucion 2254</h1><p>Norma de calidad del aire ambiente.</p></body></html>'
    text, meta = extract_text('norma.htm', html.encode())
    assert meta['format'] == 'html'
    assert 'Resolucion 2254' in text
    assert '<h1>' not in text and 'var a=1' not in text


def test_extract_docx_roundtrip():
    import docx

    buf = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph('Plan de descontaminacion del Valle de Aburra')
    doc.save(buf)
    text, meta = extract_text('plan.docx', buf.getvalue())
    assert 'descontaminacion' in text and meta['format'] == 'docx'


def test_extract_xlsx_roundtrip():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Estaciones'
    ws.append(['estacion', 'pm25'])
    ws.append(['Itagui', 42])
    buf = io.BytesIO()
    wb.save(buf)
    text, meta = extract_text('datos.xlsx', buf.getvalue())
    assert 'Itagui | 42' in text and '[hoja Estaciones]' in text


# -- embeddings ---------------------------------------------------------
def test_hash_embedder_deterministic_normed():
    e = HashEmbedder()
    assert e.dim == 384
    a = e.embed(['humo regional en el valle'])[0]
    b = e.embed(['humo regional en el valle'])[0]
    assert a == b
    assert abs(math.sqrt(sum(v * v for v in a)) - 1.0) < 1e-6
    c = e.embed(['lluvia intensa'])[0]
    assert a != c


# -- discovery (sin red) ------------------------------------------------
def test_discover_siata_fixtures():
    index = (
        '<a href="122_InformesSegundoSemestre2024/">x</a>'
        '<a href="123_InformesPrimerPeriodoGestion2025/">y</a>'
    )
    camps = discover_campaigns(index)
    assert camps[-1].endswith('123_InformesPrimerPeriodoGestion2025')
    page = (
        '<a href="33_Informe_Diario_20250327.pdf">a</a>'
        '<a href="34_Informe_Diario_20250328.pdf">b</a>'
    )
    diarios = discover_diarios(page, 'https://siata.gov.co/Informes_Aire/123_X/')
    assert diarios['20250328'].endswith('34_Informe_Diario_20250328.pdf')


@pytest.mark.asyncio
async def test_discover_follows_siata_redirects():
    """SIATA devuelve 301 sin slash: el discover debe seguirlo (sin red)."""
    import httpx
    from datetime import date

    from app.rag.sources import SIATA_INDEX, discover_siata_daily

    index_html = '<a href="123_InformesPrimerPeriodoGestion2025/">y</a>'
    camp_html = '<a href="34_Informe_Diario_20250328.pdf">a</a>'

    def handler(request):
        url = str(request.url)
        if url == SIATA_INDEX:
            return httpx.Response(301, headers={'location': SIATA_INDEX + '/'})
        if url == SIATA_INDEX + '/':
            return httpx.Response(200, text=index_html)
        if '123_InformesPrimerPeriodoGestion2025' in url:
            return httpx.Response(200, text=camp_html)
        return httpx.Response(404)

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        url = await discover_siata_daily(client, date(2025, 3, 28))
    assert url is not None and url.endswith('34_Informe_Diario_20250328.pdf')


@pytest.mark.asyncio
async def test_discover_descends_into_informes_diarios():
    """Campanas 2026+: los diarios estan en la subcarpeta InformesDiarios/ (sin red)."""
    import httpx
    from datetime import date

    from app.rag.sources import SIATA_INDEX, discover_siata_daily

    camp = '138_InformesPeriodoGestionPrimerSemestre2026'
    camp_html = '<a href="InformesDiarios/">d</a><a href="00_Otro.pdf">x</a>'
    sub_html = '<a href="34_Informe_Diario_20260915.pdf">a</a>'

    def handler(request):
        url = str(request.url)
        if url.rstrip('/') == SIATA_INDEX:
            return httpx.Response(200, text=f'<a href="{camp}/">y</a>')
        if url.endswith(camp):
            return httpx.Response(200, text=camp_html)
        if 'InformesDiarios' in url:
            return httpx.Response(200, text=sub_html)
        return httpx.Response(404)

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        url = await discover_siata_daily(client, date(2026, 9, 15))
    assert url is not None and url.endswith('34_Informe_Diario_20260915.pdf')


# -- endpoints con store falso (sin postgres) ---------------------------
class FakeStore:
    def __init__(self, embedder):
        self.embedder = embedder
        self.docs: list[dict] = []

    async def upsert_document(self, **kw):
        from app.rag.store import StoredDocument

        doc_id = f'doc-{len(self.docs) + 1}'
        vecs = self.embedder.embed([t for t, _, _ in kw['chunks']])
        self.docs.append({**kw, 'id': doc_id, 'vecs': vecs})
        return StoredDocument(
            id=doc_id, title=kw['title'], source=kw['source'],
            kind=kw['kind'], chunks=len(kw['chunks']),
        )

    async def list_documents(self):
        from app.rag.store import StoredDocument

        return [
            StoredDocument(
                id=d['id'], title=d['title'], source=d['source'],
                kind=d['kind'], chunks=len(d['chunks']),
            )
            for d in self.docs
        ]

    async def delete_document(self, doc_id: str) -> bool:
        before = len(self.docs)
        self.docs = [d for d in self.docs if d['id'] != doc_id]
        return len(self.docs) < before

    async def search(self, query, top_k=8, min_score=0.25, sources=None):
        from app.rag.store import tokens

        qv = self.embedder.embed([query])[0]
        qtok = tokens(query)
        hits: list[SearchHit] = []
        for d in self.docs:
            if sources and d['source'] not in sources:
                continue
            for i, ((t, sec, page), v) in enumerate(zip(d['chunks'], d['vecs'])):
                cos = sum(a * b for a, b in zip(qv, v))
                kw = len(qtok & tokens(t)) / max(1, len(qtok)) if qtok else 0.0
                score = 0.7 * cos + 0.3 * kw
                if score >= min_score:
                    hits.append(
                        SearchHit(
                            chunk_id=f"{d['id']}-{i}", document_id=d['id'],
                            title=d['title'], source=d['source'], kind=d['kind'],
                            uri=d.get('uri'), section=sec, page=page, text=t,
                            vector_score=round(cos, 4), keyword_score=round(kw, 4),
                            score=round(score, 4),
                        )
                    )
        hits.sort(key=lambda h: h.score, reverse=True)
        return hits[:top_k]


@pytest.fixture()
def rag_fake():
    fake = FakeStore(HashEmbedder())
    rag_ep.configure_rag(fake, HashEmbedder())
    yield fake
    rag_ep.configure_rag(None, None)


@pytest.mark.asyncio
async def test_upload_search_query_flow(rag_fake):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url='http://test') as client:
        body = (
            '# CONPES aire\nColombia adopta metas de calidad del aire. '
            'El Valle de Aburra sufre inversion termica en marzo. ' * 6
        )
        r = await client.post(
            '/api/v1/rag/documents/upload',
            files={'file': ('conpes.txt', body.encode(), 'text/plain')},
            data={'source': 'gobierno', 'kind': 'plan'},
        )
        assert r.status_code == 200, r.text
        assert r.json()['chunks'] >= 1

        r = await client.get('/api/v1/rag/documents')
        assert len(r.json()) == 1

        r = await client.post(
            '/api/v1/rag/search',
            json={'query': 'inversion termica marzo Aburra', 'min_score': 0.1},
        )
        assert r.status_code == 200
        data = r.json()
        assert data['has_source'] and len(data['hits']) >= 1
        assert data['hits'][0]['source'] == 'gobierno'

        r = await client.post('/api/v1/rag/query', json={'query': 'inversion termica marzo'})
        assert r.status_code == 200
        q = r.json()
        assert q['has_source'] and len(q['citations']) >= 1


@pytest.mark.asyncio
async def test_official_sources_registry():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url='http://test') as client:
        r = await client.get('/api/v1/rag/sources')
        assert r.status_code == 200
        names = [s['source'] for s in r.json()['sources']]
        assert {'siata', 'gobierno', 'oms', 'ods'} <= set(names)


@pytest.mark.asyncio
async def test_query_abstains_without_sources(rag_fake):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url='http://test') as client:
        r = await client.post('/api/v1/rag/query', json={'query': 'algo que no existe zzz'})
        assert r.status_code == 200
        assert r.json()['has_source'] is False


@pytest.mark.asyncio
async def test_query_chats_without_sources(monkeypatch, rag_fake):
    """Sin fuentes pero con LLM: charla que encauza, sin citas inventadas."""
    from app.core import llm as llm_mod

    async def fake_complete(messages, model=None, temperature=0.2):
        return {"content": "Hola, prueba con inversion termica.", "model": "m", "provider": "groq"}

    monkeypatch.setattr(llm_mod.omniroute, "complete", fake_complete)
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.post("/api/v1/rag/query", json={"query": "hola"})
        assert r.status_code == 200
        q = r.json()
        assert q["has_source"] is False
        assert q["provider"] == "groq"
        assert q["citations"] == []
        assert "Hola" in q["answer"]


@pytest.mark.asyncio
async def test_query_abstains_truly_without_llm(monkeypatch, rag_fake):
    """Sin fuentes y sin LLM: callejon honesto con provider abstain."""
    from app.core import llm as llm_mod

    async def boom(messages, model=None, temperature=0.2):
        raise ConnectionError("sin red")

    monkeypatch.setattr(llm_mod.omniroute, "complete", boom)
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.post("/api/v1/rag/query", json={"query": "hola"})
        assert r.status_code == 200
        q = r.json()
        assert q["has_source"] is False
        assert q["provider"] == "abstain"
        assert "No encontre fuentes" in q["answer"]


@pytest.mark.asyncio
async def test_query_extractive_fallback_cites(monkeypatch, rag_fake):
    """Sin LLM: responde con fragmentos citados [n]."""
    from app.core import llm as llm_mod

    async def boom(messages):
        raise ConnectionError('sin red')

    monkeypatch.setattr(llm_mod.omniroute, 'complete', boom)
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url='http://test') as client:
        body = '# Aire en Medellin. El PM2.5 sube en hora pico. ' * 8
        r = await client.post(
            '/api/v1/rag/documents/upload',
            files={'file': ('aire.txt', body.encode(), 'text/plain')},
            data={'source': 'proyecto', 'kind': 'otro'},
        )
        assert r.status_code == 200
        r = await client.post('/api/v1/rag/query', json={'query': 'PM2.5 hora pico'})
        assert r.status_code == 200
        q = r.json()
        assert q['has_source'] is True
        assert q['provider'] == 'extractive-fallback'
        assert '[1]' in q['answer'] and 'PM2.5' in q['answer']


@pytest.mark.asyncio
async def test_ingest_trigger_and_status(monkeypatch, rag_fake):
    """Trigger manual de ingesta SIATA + estado (sin red real)."""

    async def fake_ingest(store, embedder, day=None):
        return {"status": "ok", "day": "2026-09-16", "chunks": 3}

    monkeypatch.setattr(rag_ep, "ingest_siata_daily", fake_ingest)
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        r = await client.get("/api/v1/rag/ingest/status")
        assert r.status_code == 200
        assert r.json()["runs"] == 0
        r = await client.post("/api/v1/rag/ingest/trigger", json={})
        assert r.status_code == 200
        assert r.json()["status"] == "ok"
        r = await client.get("/api/v1/rag/ingest/status")
        assert r.json()["runs"] == 1
        assert r.json()["last_status"] == "ok"
        r = await client.post("/api/v1/rag/ingest/trigger", json={"day": "no-fecha"})
        assert r.status_code == 422


@pytest.mark.asyncio
async def test_start_scheduler_starts_and_stops():
    """El scheduler 06:00 Bogota arranca sin disparar nada en tests."""
    from app.rag.ingest import start_scheduler

    sched = start_scheduler(lambda: None, lambda: None)
    assert sched.running
    sched.shutdown(wait=False)


@pytest.mark.asyncio
async def test_upload_invalid_source_rejected(rag_fake):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url='http://test') as client:
        r = await client.post(
            '/api/v1/rag/documents/upload',
            files={'file': ('a.txt', b'hola', 'text/plain')},
            data={'source': 'nasa'},
        )
        assert r.status_code == 422