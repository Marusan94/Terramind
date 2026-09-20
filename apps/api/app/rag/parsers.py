"""Extraccion de texto real desde PDF / DOCX / XLSX / CSV / TXT / MD.

A diferencia del RAG frontend original (file.text() sobre binarios),
aqui cada formato se parsea con su libreria. Los imports son perezosos
para que el modulo cargue aunque falte una dependencia opcional.
"""
from __future__ import annotations

import csv
import io


class UnsupportedFormat(Exception):
    pass


def detect_kind(filename: str) -> str:
    ext = (filename.rsplit('.', 1)[-1] if '.' in filename else '').lower()
    return {
        'pdf': 'pdf', 'docx': 'word', 'doc': 'word',
        'xlsx': 'excel', 'xls': 'excel', 'csv': 'excel',
        'txt': 'text', 'md': 'text', 'markdown': 'text',
    }.get(ext, 'text')


def extract_text(filename: str, data: bytes) -> tuple[str, dict]:
    """Devuelve (texto, meta). Meta incluye paginas/hojas cuando aplica."""
    low = filename.lower()
    if low.endswith(('.htm', '.html', '.xhtml')):
        return _extract_html(data)
    if low.endswith('.csv'):
        return _extract_plain(data)
    kind = detect_kind(filename)
    if kind == 'pdf':
        return _extract_pdf(data)
    if kind == 'word':
        return _extract_docx(data, filename)
    if kind == 'excel':
        return _extract_sheet(data, filename)
    return _extract_plain(data)


def _extract_plain(data: bytes) -> tuple[str, dict]:
    text = data.decode('utf-8', errors='replace')
    if text.lstrip().startswith(('{', '[')):
        return text, {'format': 'text'}
    # CSV: filas legibles "celda | celda".
    try:
        rows = list(csv.reader(io.StringIO(text)))
        if rows and any(len(r) > 1 for r in rows):
            text = '\n'.join(' | '.join(c.strip() for c in r) for r in rows if any(c.strip() for c in r))
            return text, {'format': 'csv', 'rows': len(rows)}
    except csv.Error:
        pass
    return text, {'format': 'text'}


def _extract_pdf(data: bytes) -> tuple[str, dict]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise UnsupportedFormat('pypdf no instalado: pip install pypdf') from exc
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            t = page.extract_text() or ''
        except Exception:
            t = ''
        t = t.replace('\x00', ' ').strip()
        if t:
            pages.append(f'[pag {i}]\n{t}')
    text = '\n\n'.join(pages).strip()
    # Deshifeniza cortes de linea tipo "contamina-\ncion".
    text = text.replace('-\n', '')
    return text, {'format': 'pdf', 'pages': len(reader.pages), 'pages_with_text': len(pages)}


def _extract_docx(data: bytes, filename: str) -> tuple[str, dict]:
    if filename.lower().endswith('.doc'):
        raise UnsupportedFormat('.doc legacy no soportado, convierte a .docx')
    try:
        import docx
    except ImportError as exc:
        raise UnsupportedFormat('python-docx no instalado: pip install python-docx') from exc
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts).strip(), {'format': 'docx', 'paragraphs': len(parts)}


def _extract_sheet(data: bytes, filename: str) -> tuple[str, dict]:
    if filename.lower().endswith('.xls'):
        raise UnsupportedFormat('.xls legacy no soportado, convierte a .xlsx')
    try:
        import openpyxl
    except ImportError as exc:
        raise UnsupportedFormat('openpyxl no instalado: pip install openpyxl') from exc
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[str] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip() != '']
            if cells:
                rows.append(' | '.join(cells))
        if rows:
            blocks.append(f'[hoja {ws.title}]\n' + '\n'.join(rows))
    return '\n\n'.join(blocks).strip(), {'format': 'xlsx', 'sheets': len(wb.worksheets)}

def _extract_html(data: bytes) -> tuple[str, dict]:
    """Texto legible desde paginas .htm/.html (portales oficiales)."""
    import html as _html
    from html.parser import HTMLParser

    class _TextOnly(HTMLParser):
        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.parts: list[str] = []
            self._skip = False

        def handle_starttag(self, tag: str, attrs) -> None:
            if tag in ('script', 'style', 'nav', 'header', 'footer'):
                self._skip = True
            elif tag in ('p', 'br', 'div', 'li', 'tr', 'h1', 'h2', 'h3', 'h4'):
                self.parts.append(chr(10))

        def handle_endtag(self, tag: str) -> None:
            if tag in ('script', 'style', 'nav', 'header', 'footer'):
                self._skip = False

        def handle_data(self, data: str) -> None:
            if not self._skip:
                text = ' '.join(data.split())
                if text:
                    self.parts.append(text)

    parser = _TextOnly()
    parser.feed(data.decode('utf-8', errors='replace'))
    text = chr(10).join(line.strip() for line in chr(10).join(parser.parts).split(chr(10)))
    text = _html.unescape(chr(10).join(line for line in text.split(chr(10)) if line))
    return text.strip(), {'format': 'html'}